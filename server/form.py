import os
import re
from docx import Document
import pandas as pd
from tkinter import Tk, TclError
from tkinter import ttk  # Themed tkinter
import argparse

def extract_placeholders_from_docx(filepath):
    placeholders = set()
    pattern = re.compile(r'\{\{\s*(.+?)\s*\}\}')

    try:
        document = Document(filepath)
        for para in document.paragraphs:
            found = pattern.findall(para.text)
            placeholders.update(found)
    except Exception as e:
        print(f"Error processing file {filepath}: {e}")
    return placeholders

def collect_data_from_user(placeholders, root, output_path):
    entries = {}
    frame = ttk.Frame(root)
    frame.pack(pady=20, padx=20)

    style = ttk.Style()
    style.configure("TLabel", foreground="navy", font=("Helvetica", 12), padding=10)
    style.configure("TEntry", foreground="black", padding=10)
    # Custom button style
    style.configure("TButton", foreground="white", background="blue", font=("Helvetica", 12), padding=10)

    for i, field in enumerate(placeholders):
        ttk.Label(frame, text=field, style="TLabel").grid(row=i, column=0, sticky='e')
        entry = ttk.Entry(frame, style="TEntry")
        entry.grid(row=i, column=1, sticky='we')
        entries[field] = entry

    def add_entry():
        data = {field: entry.get() for field, entry in entries.items()}
        save_data_to_excel(data, filename=os.path.join(output_path, 'data.xlsx'))
        for entry in entries.values():
            entry.delete(0, 'end')

    def submit_and_close():
        add_entry()  # Save the current data first
        root.quit()  # Close the application

    ttk.Button(frame, text='Add Another Entry', command=add_entry, style="TButton").grid(row=len(placeholders) + 1, column=0)
    ttk.Button(frame, text='Submit and Close', command=submit_and_close, style="TButton").grid(row=len(placeholders) + 1, column=1)

def save_data_to_excel(data, filename='data.xlsx'):
    if os.path.exists(filename):
        df = pd.read_excel(filename, index_col=None)
    else:
        df = pd.DataFrame()

    new_df = pd.DataFrame([data])
    df = pd.concat([df, new_df], ignore_index=True)
    df.to_excel(filename, index=False)
    print("Data saved to Excel successfully.")

def main(directory, output_path):
    if not os.path.exists(output_path):
        os.makedirs(output_path)  # Create output directory if it does not exist

    docx_files = []
    if os.path.isdir(directory):
        docx_files = [os.path.join(directory, f) for f in os.listdir(directory) if f.endswith('.docx')]

    combined_placeholders = set()
    for docx_path in docx_files:
        placeholders = extract_placeholders_from_docx(docx_path)
        combined_placeholders.update(placeholders)

    if combined_placeholders:
        root = Tk()
        root.title("Data Collection Form")
        style = ttk.Style(root)
        # Setting a theme that allows for background color changes on buttons
        try:
            style.theme_use('clam')  # Clam theme usually allows more flexibility for background colors
        except TclError:
            pass
        collect_data_from_user(combined_placeholders, root, output_path)
        root.mainloop()
    else:
        print("No placeholders found to process.")

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description='Process some DOCX files.')
    parser.add_argument('--docx_directory', type=str, help='The directory where DOCX files are stored')
    parser.add_argument('--output_directory', type=str, help='The directory where the Excel file will be saved')
    args = parser.parse_args()

    main(args.docx_directory, args.output_directory)
